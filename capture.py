#!/usr/bin/env python3
"""
Capture Grafana panels and insert them into a pre-made Excel template.
"""

import json
import os
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright

TZ = ZoneInfo("Asia/Bangkok")

SLOT_BY_CRON = {
    "0 0 * * *": "07:00",
    "0 1 * * *": "08:00",
    "0 2 * * *": "09:00",
    "0 5 * * *": "12:00",
    "0 6 * * *": "13:00",
    "0 12 * * *": "19:00",
    "0 13 * * *": "20:00",
}

ROOT = Path(__file__).resolve().parent
CFG = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
NOW = datetime.now(TZ)


def resolve_slot() -> str:
    slot = os.getenv("SLOT", "").strip()
    schedule = os.getenv("SCHEDULE", "").strip()

    if not slot:
        slot = SLOT_BY_CRON.get(schedule)

    if slot not in CFG["anchors"]:
        print(
            f"ERROR: unknown slot {slot!r}; "
            f"known={list(CFG['anchors'])}; "
            f"SLOT={os.getenv('SLOT')!r}; "
            f"SCHEDULE={schedule!r}",
            file=sys.stderr,
        )
        sys.exit(1)

    return slot


def anchor_of(img) -> str:
    anchor = img.anchor

    if isinstance(anchor, str):
        return anchor

    return (
        f"{get_column_letter(anchor._from.col + 1)}"
        f"{anchor._from.row + 1}"
    )


def drop_old_image(ws, cell: str) -> None:
    """Remove an image already anchored at the specified Excel cell."""
    for img in list(ws._images):
        if anchor_of(img) == cell:
            ws._images.remove(img)


def ensure_logged_in(page, metric: str) -> None:
    current_url = page.url.lower()

    if "/login" in current_url:
        raise RuntimeError(
            f"Grafana session expired while opening {metric}: {page.url}"
        )


def main() -> None:
    slot = resolve_slot()
    date = NOW.strftime("%Y-%m-%d")

    reports_dir = ROOT / "reports"
    shots_dir = ROOT / "screenshots" / date / slot.replace(":", "")

    reports_dir.mkdir(parents=True, exist_ok=True)
    shots_dir.mkdir(parents=True, exist_ok=True)

    report_path = reports_dir / f"{date}.xlsx"
    template_path = ROOT / "template.xlsx"

    source = report_path if report_path.exists() else template_path

    if not source.exists():
        print(f"ERROR: Excel source not found: {source}", file=sys.stderr)
        sys.exit(1)

    state = ROOT / "storage_state.json"

    if not state.exists():
        print(
            "ERROR: storage_state.json missing "
            "(restore it from GRAFANA_STORAGE_STATE_B64)",
            file=sys.stderr,
        )
        sys.exit(1)

    wb = load_workbook(source)

    sheet_name = CFG["sheet"]
    if sheet_name not in wb.sheetnames:
        print(
            f"ERROR: worksheet {sheet_name!r} not found. "
            f"Available sheets: {wb.sheetnames}",
            file=sys.stderr,
        )
        sys.exit(1)

    ws = wb[sheet_name]
    anchors = CFG["anchors"][slot]
    failures = []

    with sync_playwright() as p:
        with p.chromium.launch(headless=True) as browser:
            ctx = browser.new_context(
                storage_state=str(state),
                viewport={"width": 1600, "height": 1000},
                locale="en-US",
                # Required temporarily because the internal Grafana
                # certificate has a hostname/CA mismatch.
                ignore_https_errors=True,
            )

            page = ctx.new_page()
            page.set_default_timeout(60_000)
            page.set_default_navigation_timeout(60_000)

            for metric, url in CFG["urls"].items():
                if metric not in anchors:
                    print(f"Skipping {metric}: no Excel anchor configured")
                    continue

                print(f"[{slot}] capturing {metric} ...")

                try:
                    page.goto(
                        url,
                        wait_until="domcontentloaded",
                        timeout=60_000,
                    )

                    # Check immediately after navigation.
                    ensure_logged_in(page, metric)

                    # Allow Grafana panels and charts to render.
                    page.wait_for_timeout(8_000)

                    # Check again in case Grafana redirected after loading.
                    ensure_logged_in(page, metric)

                    png = shots_dir / f"{metric}.png"
                    page.screenshot(
                        path=str(png),
                        full_page=True,
                    )

                    cell = anchors[metric]
                    drop_old_image(ws, cell)

                    img = XLImage(str(png))
                    img.width = 520
                    img.height = 300
                    ws.add_image(img, cell)

                    print(f"[{slot}] OK: {metric} -> {png}")

                except (PlaywrightTimeoutError, Exception) as exc:
                    message = f"{metric}: {type(exc).__name__}: {exc}"
                    print(f"ERROR: {message}", file=sys.stderr)
                    failures.append(message)

            ctx.close()

    if failures:
        print("\nCapture failed for one or more metrics:", file=sys.stderr)
        for failure in failures:
            print(f"- {failure}", file=sys.stderr)

        # Do not save/commit a partially updated workbook.
        sys.exit(2)

    ws[CFG["date_cell"]] = NOW.strftime("%d %b %Y")
    wb.save(report_path)

    print(f"OK -> {report_path}")


if __name__ == "__main__":
    main()
